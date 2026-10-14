#!/usr/bin/env python3
"""Lager et gjennomgangs-regneark (.xlsx) fra data/kar-icd10-synonymer.json.

Bruk:  python3 tools/lag_regneark.py UTFIL.xlsx
Regnearket åpnes i Excel/Numbers. Brukeren redigerer kolonnene Synonymer,
Instrukser og Ferdig?, og sender filen tilbake for implementering i JSON.
Kolonnen «Opprinnelig» brukes til å se hva som er endret og skal ikke redigeres.
"""
import json, sys, pathlib, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

root = pathlib.Path(__file__).resolve().parent.parent
data = json.loads((root / "data" / "kar-icd10-synonymer.json").read_text(encoding="utf8"))
out = sys.argv[1] if len(sys.argv) > 1 else "karma-diagnose-synonymer.xlsx"
SEP = "; "

# ---- Claude sine forslag (korte ord, siden OG-søket også ser kodeteksten) ----
FORSLAG = {
 "G45.1": ["TIA carotis", "forbigående lammelse", "forbigående afasi", "forbigående talevansker", "forbigående hemiparese"],
 "G45.3": ["forbigående blindhet", "synstap forbigående", "gardin foran øyet", "monokulær blindhet", "retinal TIA"],
 "I70.0": ["aortaskleross", "forkalket aorta", "kalk i aorta", "aortaplakk"],
 "I70.00": ["aortaskleross", "forkalket aorta", "kalk i aorta", "aortaplakk"],
 "I70.01": ["aortaskleross", "forkalket aorta", "kalk i aorta", "aortaplakk"],
 "I70.1": ["renovaskulær hypertensjon", "stenose nyrearterie", "nyrearterieaterosklerose"],
 "I70.10": ["renovaskulær hypertensjon", "stenose nyrearterie", "nyrearterieaterosklerose"],
 "I70.11": ["renovaskulær hypertensjon", "stenose nyrearterie", "nyrearterieaterosklerose"],
 "I70.2": ["PAVK", "perifer arteriell karsykdom", "aterosklerose bein"],
 "I70.20": ["PAVK", "SFA-okklusjon", "SFA-stenose", "iliacaokklusjon", "iliacastenose", "kort gangdistanse", "redusert gangdistanse", "leggsmerter ved gange", "trange kar i bena"],
 "I70.21": ["PAVK med gangren", "iskemisk fotsår", "tåsår", "hælsår", "tørr gangren", "våt gangren", "nekrose tå", "nekrose fot", "iskemisk sår"],
 "I70.8": ["mesenterial aterosklerose", "stenose a. mesenterica", "stenose truncus coeliacus", "visceral aterosklerose", "angina abdominalis"],
 "I70.80": ["mesenterial aterosklerose", "stenose a. mesenterica", "stenose truncus coeliacus", "visceral aterosklerose", "angina abdominalis"],
 "I72.1": ["subclavia", "subclaviaaneurisme", "axillaris", "armarterie"],
 "I72.3": ["iliaca", "iliacaaneurisme", "bekkenarterie"],
 "I72.5": ["carotisaneurisme", "vertebralisaneurisme", "halsarterie"],
 "I72.8": ["visceral", "a. lienalis", "miltarterie", "a. renalis", "nyrearterie", "a. mesenterica", "a. hepatica", "a. coeliaca"],
 "I73.8": ["erytromelalgi"],
 "I74.1": ["aortaemboli", "thorakal aortatrombose"],
 "I74.4": ["akutt ekstremitetsiskemi", "embolus ekstremitet"],
 "I74.9": ["arteriell emboli", "arteriell trombose", "embolus"],
 "I77.1": ["carotis coiling"],
 "I77.2": ["traumatisk karskade"],
 "I77.6": ["storkarsvaskulitt"],
 "I82.1": ["migrerende tromboflebitt", "Trousseau"],
 "I82.8": ["trombose v. subclavia", "Paget-Schroetter", "effort-trombose", "armvenetrombose", "trombose v. jugularis", "trombose v. axillaris"],
 "I89.1": ["lymfebetennelse"],
 "I89.8": ["lymfocele", "lymfelekkasje", "lymfefistel", "lymfocele lyske"],
 "T81.0": ["pseudoaneurisme lyske"],
 "T81.1": ["postoperativt sjokk"],
 "T81.5": ["glemt kompress", "fremmedlegeme etter operasjon"],
 "T87.0": ["replantasjon arm", "replantasjon hånd", "replantasjon finger"],
 "T87.1": ["replantasjon ben", "replantasjon fot"],
 "L97": ["leggsår", "arterielt sår", "iskemisk sår", "venøst leggsår", "kronisk sår bein"],
 "N18.1": ["CKD 1"], "N18.2": ["CKD 2"],
 "N18.3": ["CKD 3", "nyresvikt stadium 3"], "N18.4": ["CKD 4", "nyresvikt stadium 4"],
 "Q25.4": ["medfødt aortabue-anomali"],
}
for c in ("S15.2", "S15.3", "S15.7", "S15.8"): FORSLAG[c] = ["halskarskade"]
for c in ("S65.0", "S65.1", "S65.2", "S65.3", "S65.7", "S65.8", "S65.9"): FORSLAG[c] = ["håndkarskade"]
for c in ("S65.4", "S65.5"): FORSLAG[c] = ["håndkarskade", "fingerkarskade"]
for c in ("S75.0", "S75.1", "S75.2", "S75.7", "S75.8", "S75.9"): FORSLAG[c] = ["lårkarskade"]
for c in ("S95.0", "S95.1", "S95.2", "S95.7", "S95.8", "S95.9"): FORSLAG[c] = ["fotkarskade", "ankelkarskade"]

MERKNAD = {
 "I70.0": "Skal «med gangren»-undergruppene (.00/.01, .10/.11, .80/.81, .90/.91) ha egne synonymer, eller bare arve overordnet kode? Gangren i aorta/nyrearterie er sjeldent.",
 "I70.20": "Sår uten gangren (arterielt/iskemisk sår): skal det peke hit med L97 som bidiagnose, eller til I70.21? Hvilesmerter ligger i dag her.",
 "I70.21": "Skal «hvilesmerter»/«rest pain» også foreslå I70.21 i tillegg til I70.20?",
 "I70.8": "Hva bruker dere I70.8 til i praksis (mesenterial/visceral)? Bekreft forslagene.",
 "I65.8": "Hvilken kode bruker dere for stenose i a. subclavia / truncus brachiocephalicus?",
 "I74.3": "Ord som «emboli», «embolus» og «akutt iskemi» finnes i mange koder (I74.x, I82.x, T81.x …). Skal de rangere I74.3 øverst? Skriv rekkefølgen i Instrukser.",
 "I74.9": "«emboli» alene er generisk – bør det vektes lavere enn I74.3?",
}

DISPLAY_GRUPPE = {
 "G45 Forbigaaende cerebrale iskemiske anfall": "G45 Forbigående cerebrale iskemiske anfall",
 "G46 Cerebrale karsyndromer ved cerebrovaskulaer sykdom": "G46 Cerebrale karsyndromer ved cerebrovaskulær sykdom",
 "I71 Aortaaneurisme og -disseksjon": "I71 Aortaaneurisme og -disseksjon",
 "I78 Sykdommer i kapillarer": "I78 Sykdommer i kapillærer",
 "I82 Annen venoes emboli og trombose": "I82 Annen venøs emboli og trombose",
 "I83 Aareknuter i underekstremitet": "I83 Åreknuter i underekstremitet",
 "I85 Oesofagusvaricer": "I85 Øsofagusvaricer",
 "I86 Aareknuter i andre lokalisasjoner": "I86 Åreknuter i andre lokalisasjoner",
 "I89 Andre ikke-infektioese sykdommer i lymfekar/lymfeknuter": "I89 Andre ikke-infeksiøse sykdommer i lymfekar/lymfeknuter",
 "Karskader (traume) - haandledd/haand": "Karskader (traume) – håndledd/hånd",
 "Karskader (traume) - hofte/laar": "Karskader (traume) – hofte/lår",
 "Leggsaar": "Leggsår",
 "Medfoedte karmisdannelser": "Medfødte karmisdannelser",
 "T82 Komplikasjoner av kardielle og vaskulere proteser, implantater og transplantater": "T82 Komplikasjoner av kardielle og vaskulære proteser, implantater og transplantater",
}
def gnavn(g): return DISPLAY_GRUPPE.get(g, g.replace(" - ", " – "))

# ---- Bygg rader ----
def rows_for(lst):
    res = []
    for c in lst:
        orig = list(c.get("synonyms", []))
        assert all(";" not in s for s in orig), c["code"]
        have = {s.lower() for s in orig}
        add = [s for s in FORSLAG.get(c["code"], []) if s.lower() not in have]
        res.append(dict(code=c["code"], tekst=c["term_no"], gruppe=c["gruppe"], rel=c["relevans"],
                        orig=orig, ny=orig + add, lagt_til=bool(add),
                        bi=c.get("bidiagnose_forslag") or [], merk=MERKNAD.get(c["code"], "")))
    return res
main_rows, adj_rows = rows_for(data["codes"]), rows_for(data["tilstotende_koder"])
all_codes = {r["code"] for r in main_rows + adj_rows}
assert set(FORSLAG) <= all_codes, set(FORSLAG) - all_codes

# ---- Stiler ----
F = "Arial"
f_norm, f_bold = Font(name=F, size=10), Font(name=F, size=10, bold=True)
f_hdr = Font(name=F, size=10, bold=True, color="FFFFFF")
f_grey = Font(name=F, size=9, color="808080")
fill_hdr = PatternFill("solid", fgColor="0F6B63")
fill_sec = PatternFill("solid", fgColor="0A4F49")
fill_grp = PatternFill("solid", fgColor="D5E8E5")
fill_edit = PatternFill("solid", fgColor="FFFFFF")
fill_new = PatternFill("solid", fgColor="FFF2B3")
fill_note = PatternFill("solid", fgColor="FBF0E0")
thin = Side(style="thin", color="DCE6E3"); box = Border(top=thin, bottom=thin, left=thin, right=thin)
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()

# ---- Ark 1: Les først ----
ws0 = wb.active; ws0.title = "Les først"
ws0.sheet_view.showGridLines = False
ws0.column_dimensions["A"].width = 4; ws0.column_dimensions["B"].width = 110
n_ny = sum(r["lagt_til"] for r in main_rows + adj_rows)
lines = [
 ("Karma Diagnose – gjennomgang av synonymer", "title"),
 (f"Laget {datetime.date.today().isoformat()} fra data/kar-icd10-synonymer.json ({len(main_rows)} hovedkoder + {len(adj_rows)} tilstøtende koder).", "grey"),
 ("", None),
 ("Slik jobber du", "h"),
 ("1. Gå til arket «Synonymer». Kodene står i samme rekkefølge og hierarki som ICD-10 (gruppe → kode → undergruppe, innrykk viser nivå).", None),
 ("2. Kolonne C «Synonymer» er den du redigerer. Ett synonym per bit, atskilt med semikolon og mellomrom:  PAVK; claudicatio; røykebein", None),
 ("   Gul celle = jeg har lagt til nye forslag. Se kolonne I «Opprinnelig» for hva som lå i appen før. Slett det du ikke vil ha, legg til det du mangler.", None),
 ("3. Kolonne D «Dine instrukser» er fritekst om vekting, rangering og bidiagnoser (eksempler under). Skriv slik du ville forklart det til en kollega.", None),
 ("4. Kolonne E «Merknad fra Claude» er spørsmål jeg trenger svar på. Svar i kolonne D på samme rad.", None),
 ("5. Sett «Ja» i kolonne F når raden er gjennomgått. Rader uten «Ja» behandler jeg som uendret, så du kan ta det gradvis og sende filen flere ganger.", None),
 ("6. Ikke endre kolonne A (kode) eller kolonne I (Opprinnelig). Ikke slett eller flytt rader. Tomme rader og gruppeoverskrifter ignoreres.", None),
 ("7. Send filen tilbake til meg. Jeg legger endringene inn i JSON-filen, bygger appen på nytt og tester før det pushes til dev.", None),
 ("", None),
 ("Eksempler på instrukser (kolonne D)", "h"),
 ("• «emboli» (og «embolus») skal gi I74.3 øverst, deretter I74.2, så de øvrige I74.x.", None),
 ("• «AAA» skal alltid ha I71.4 øverst, I71.3 hvis ordet «ruptur» eller «rumpert» også er med.", None),
 ("• Ved I70.21: foreslå L97 som bidiagnose.", None),
 ("• «slag» alene er for generelt – gi lav vekt, men «hjerneslag» skal treffe I63.9 øverst.", None),
 ("• Dette synonymet hører ikke hjemme her, men under I80.2.", None),
 ("", None),
 ("Gode å vite", "h"),
 ("• Søket krever først at alle ord finnes i kodeteksten + synonymene samlet (OG), og faller tilbake til minst ett ord (ELLER). Derfor holder det med korte ord («poplitea») i stedet for hele fraser («aneurisme a. poplitea»).", None),
 ("• Stavevarianter som aa/å, oe/ø og tromb/thromb, samt aterosklerose/arteriosklerose, håndteres av programmet. Du trenger ikke skrive begge.", None),
 ("• Arket «Eksisterende regler» viser termgrupper, søkeregler og prioriteringer som allerede finnes, så du ikke gjentar dem.", None),
 (f"• Jeg har lagt til forslag på {n_ny} rader. Alle er ment som utkast – kast gjerne det som ikke stemmer med hvordan dere faktisk skriver.", None),
]
for i, (t, kind) in enumerate(lines, start=2):
    c = ws0.cell(row=i, column=2, value=t); c.alignment = Alignment(wrap_text=True, vertical="top")
    c.font = {"title": Font(name=F, size=14, bold=True, color="0F6B63"), "h": Font(name=F, size=11, bold=True, color="0F6B63"),
              "grey": f_grey}.get(kind, f_norm)

# ---- Ark 2: Synonymer ----
ws = wb.create_sheet("Synonymer")
headers = ["Kode", "Offisiell tekst (Helsedirektoratet)", "Synonymer (rediger her)", "Dine instrukser (vekting, rangering, bidiagnose …)",
           "Merknad fra Claude", "Ferdig?", "Relevans", "Bidiagnose i dag", "Opprinnelig (ikke rediger)"]
widths = [13, 46, 56, 44, 40, 8, 10, 14, 44]
for j, (h, w) in enumerate(zip(headers, widths), start=1):
    c = ws.cell(row=1, column=j, value=h); c.font = f_hdr; c.fill = fill_hdr; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = box
    ws.column_dimensions[chr(64 + j)].width = w
ws.row_dimensions[1].height = 32
ws.freeze_panes = "C2"
r = 2
def section(title):
    global r
    for j in range(1, 10):
        c = ws.cell(row=r, column=j); c.fill = fill_sec
    ws.cell(row=r, column=1, value=title).font = f_hdr
    r += 1
def build(rows):
    global r
    codes = {x["code"] for x in rows}; last = None
    for x in rows:
        if x["gruppe"] != last:
            for j in range(1, 10): ws.cell(row=r, column=j).fill = fill_grp
            ws.cell(row=r, column=1, value=gnavn(x["gruppe"])).font = f_bold
            r += 1; last = x["gruppe"]
        level = max(0, len(x["code"].replace(".", "")) - 3)
        has_children = any(o != x["code"] and o.startswith(x["code"]) for o in codes)
        vals = [x["code"], x["tekst"], SEP.join(x["ny"]), "", x["merk"], "", x["rel"].replace("hoy", "høy"),
                ", ".join(b["kode"] if isinstance(b, dict) else str(b) for b in x["bi"]), SEP.join(x["orig"])]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v); c.font = f_norm; c.alignment = wrap; c.border = box
        ws.cell(row=r, column=1).alignment = Alignment(indent=level * 2, vertical="top")
        ws.cell(row=r, column=1).font = f_bold if has_children else f_norm
        ws.cell(row=r, column=2).font = f_bold if has_children else f_norm
        if x["lagt_til"]: ws.cell(row=r, column=3).fill = fill_new
        if x["merk"]: ws.cell(row=r, column=5).fill = fill_note
        ws.cell(row=r, column=6).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=r, column=9).font = f_grey
        r += 1
section("HOVEDKODER"); build(main_rows)
section("TILSTØTENDE KODER (diabetes, nyresvikt, tarmkar, sår, traume, …)"); build(adj_rows)
dv = DataValidation(type="list", formula1='"Ja"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"F2:F{r}")
ws.auto_filter.ref = f"A1:I{r - 1}"

# ---- Ark 3: Eksisterende regler ----
w3 = wb.create_sheet("Eksisterende regler")
for col, w in zip("ABCD", (30, 70, 55, 50)): w3.column_dimensions[col].width = w
rr = 1
def head(t, cols):
    global rr
    w3.cell(row=rr, column=1, value=t).font = Font(name=F, size=12, bold=True, color="0F6B63"); rr += 1
    for j, h in enumerate(cols, start=1):
        c = w3.cell(row=rr, column=j, value=h); c.font = f_hdr; c.fill = fill_hdr
    rr += 1
def row(*vals):
    global rr
    for j, v in enumerate(vals, start=1):
        c = w3.cell(row=rr, column=j, value=v); c.font = f_norm; c.alignment = wrap
    rr += 1
head("Termgrupper (ord som alltid behandles som det samme)", ["Navn", "Kanonisk ord + synonymer", "Kommentar"])
for t in data["termgrupper"]: row(t["id"], SEP.join([t["kanonisk"]] + t["synonymer"]), t.get("kommentar") or "")
rr += 1
head("Søkeregler (hvilke koder som vises, og i hvilken rekkefølge)", ["Navn", "Beskrivelse", "Utløses av", "Resultat (kode → bidiagnose, vekt)"])
for s in data["sokeregler"]:
    res = "; ".join(f'{x["kode"]}' + (f' (bidiagnose {x["bidiagnose"]})' if x.get("bidiagnose") else "") + f', vekt {x["vekt"]}' for x in s["resultat"])
    trig = s["trigger_termgruppe"] + (" + ordet: " + " / ".join(s["krever_modifier_ord"]) if s.get("krever_modifier_ord") else "")
    row(s["id"], s["beskrivelse"], trig, res)
rr += 1
head("Prioritering (koder som skal ligge øverst for bestemte ord)", ["Navn", "Beskrivelse", "Søkeord", "Kode og rekkefølge"])
for p in data.get("prioriteringsregler", []):
    row(p["id"], p["beskrivelse"], SEP.join(p["trigger_ord"]), "; ".join(f'{x["kode"]} (rekkefølge {x["vekt"]})' for x in p["prioriteringer"]))
rr += 1
head("Stavenormalisering (gjelder automatisk)", ["Type", "Ord", "Kommentar"])
for n in data["normaliseringsregler"]:
    row(n.get("type", ""), SEP.join(n.get("par") or n.get("termer") or []), (n.get("kommentar") or "")[:300])

wb.save(out)
print("lagret", out, "| rader med forslag:", n_ny, "| totalt rader:", r - 2)
